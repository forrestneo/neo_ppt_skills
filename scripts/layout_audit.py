# -*- coding: utf-8 -*-
"""layout_audit — 成品 PPTX 静态排版审计（不渲染即可判定）。

移植自 pptx-harness（MIT）src/pptx_harness/audit.py；阈值口径与
references/stylekit/layout-engine.md 一致。对任意成品（模式 B 还原稿 /
模式 D 出稿 / compose 组装稿）做交付前门禁：

  out-of-bounds    形状超出画布 >1px 容差（error）
  placeholder      残留占位符文本（error）
  dangling-rel     图片 rId 无对应关系 / 关系指向空部件（error）
  overlap          文本形状两两包围盒重叠 >较小盒面积 10%（warning）
  overflow-est     启发式溢出：CJK 全角 1.0em、其余 0.53em、行高 1.25，
                   估算高 >可用高 105%（warning）
  orphan-line      中文孤字：折行后末行仅 ~1 个全角字宽（warning，修复靠
                   重排行宽或删非必要措辞，勿缩全段字号）
  font-ea-missing  run 设了 latin 却无 ea 且含中文（warning，中文最大坑）

用法:
  python scripts/layout_audit.py 成品.pptx                 # 全量检查
  python scripts/layout_audit.py 成品.pptx --level error   # 只看 error
  python scripts/layout_audit.py 成品.pptx --json out.json # 结构化输出

退出码: 0 = 无 error（warning 仅报告）；1 = 有 error；2 = 用法/文件错误。
"""
from __future__ import annotations

import json
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

_PLACEHOLDER_PAT = re.compile(r"click to add|lorem|\[insert|\bxxx+\b|\btodo\b", re.I)
_EM = 12700  # 1pt = 12700 EMU


@dataclass
class Finding:
    severity: str      # error | warning
    rule: str
    slide: int         # 1-based
    shape_id: int | None
    message: str


def _walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                yield from _walk(sh.shapes)
            except Exception:
                pass


def _para_font_pt(para):
    for r in para.runs:
        try:
            if r.font.size is not None:
                return r.font.size.pt
        except Exception:
            pass
    try:
        if para.font.size is not None:
            return para.font.size.pt
    except Exception:
        pass
    return None


def _est_width_emu(text: str, pt: float) -> int:
    units = sum(1.0 if ord(c) > 0x2E80 else 0.53 for c in text)
    return int(units * pt * _EM)


def _text_findings(shape, slide_no, out):
    tf = shape.text_frame
    if not tf.text.strip():
        return
    if _PLACEHOLDER_PAT.search(tf.text):
        out.append(Finding("error", "placeholder", slide_no, shape.shape_id,
                           f"残留占位符: {tf.text.strip()[:40]!r}"))
    w, h = shape.width, shape.height
    if not w or not h:
        return
    try:
        ml, mr = tf.margin_left or 0, tf.margin_right or 0
        mt, mb = tf.margin_top or 0, tf.margin_bottom or 0
    except Exception:
        ml = mr = 91440
        mt = mb = 45720
    avail_w = max(w - ml - mr, 1)
    avail_h = max(h - mt - mb, 1)
    total_h = 0
    for para in tf.paragraphs:
        pt = _para_font_pt(para) or 18.0
        if not para.text:
            total_h += int(1.25 * pt * _EM)
            continue
        lines = max(1, -(-_est_width_emu(para.text, pt) // avail_w))
        ls = 1.25
        try:
            if isinstance(para.line_spacing, float) and para.line_spacing > 0:
                ls = para.line_spacing
        except Exception:
            pass
        total_h += int(lines * ls * pt * _EM)
        if lines >= 2:
            last_w = _est_width_emu(para.text, pt) - (lines - 1) * avail_w
            if last_w <= int(1.1 * pt * _EM):
                out.append(Finding(
                    "warning", "orphan-line", slide_no, shape.shape_id,
                    f"末行孤字: …{para.text[-6:]!r}（重排行宽或删非必要措辞）"))
        for r in para.runs:
            try:
                rPr = r._r.rPr
                if (rPr is not None and rPr.find(qn("a:latin")) is not None
                        and rPr.find(qn("a:ea")) is None
                        and any(ord(c) > 0x2E80 for c in (r.text or ""))):
                    out.append(Finding(
                        "warning", "font-ea-missing", slide_no, shape.shape_id,
                        f"run 设了 latin 而无 ea（中文可能不生效）: {(r.text or '')[:20]!r}"))
                    break
            except Exception:
                pass
    if total_h > avail_h * 1.05:
        out.append(Finding(
            "warning", "overflow-est", slide_no, shape.shape_id,
            f"估算文本高 {total_h / _EM / 72:.2f}in > 盒高 {avail_h / _EM / 72:.2f}in"))


def _bounds_findings(shape, slide_no, sw, sh_h, out):
    try:
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return
    if l is None or t is None or w is None or h is None:
        return
    tol = 9525  # 1px
    if l < -tol or t < -tol or l + w > sw + tol or t + h > sh_h + tol:
        out.append(Finding(
            "error", "out-of-bounds", slide_no, shape.shape_id,
            f"形状超出画布: L{l / _EM / 72:.2f} T{t / _EM / 72:.2f} "
            f"R{(l + w) / _EM / 72:.2f} B{(t + h) / _EM / 72:.2f} (in)"))


def _overlap_findings(shapes, slide_no, out):
    boxes = []
    for sh in shapes:
        try:
            if sh.has_text_frame and sh.text_frame.text.strip() and sh.width and sh.height:
                boxes.append((sh, sh.left, sh.top, sh.width, sh.height))
        except Exception:
            continue
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ix = max(0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]))
            iy = max(0, min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2]))
            inter = ix * iy
            smaller = min(a[3] * a[4], b[3] * b[4])
            if smaller and inter / smaller > 0.10:
                out.append(Finding(
                    "warning", "overlap", slide_no, a[0].shape_id,
                    f"文本形状重叠 {inter / smaller:.0%}（与 shape {b[0].shape_id}）"))


def _rel_findings(slide, slide_no, out):
    rels = slide.part.rels
    for el in slide._element.iter(qn("a:blip")):
        rid = el.get(qn("r:embed")) or el.get(qn("r:link"))
        if rid and rid not in rels:
            out.append(Finding("error", "dangling-rel", slide_no, None,
                               f"图片 rId {rid} 无对应关系"))
    for rid, rel in rels.items():
        if rel.is_external:
            continue
        try:
            if not rel.target_part.blob:
                out.append(Finding("error", "empty-part", slide_no, None,
                                   f"关系 {rid} 指向空部件"))
        except Exception as e:
            out.append(Finding("error", "dangling-rel", slide_no, None,
                               f"关系 {rid} 解析失败: {e}"))


def audit(prs, level: str = "warning") -> list:
    findings = []
    sw, sh_h = prs.slide_width, prs.slide_height
    for i, slide in enumerate(prs.slides, 1):
        shapes = list(_walk(slide.shapes))
        for sh in shapes:
            _bounds_findings(sh, i, sw, sh_h, findings)
            try:
                if sh.has_text_frame:
                    _text_findings(sh, i, findings)
            except Exception:
                pass
        _overlap_findings(shapes, i, findings)
        _rel_findings(slide, i, findings)
    order = {"error": 0, "warning": 1}
    findings.sort(key=lambda f: (f.slide, order.get(f.severity, 9)))
    if level == "error":
        findings = [f for f in findings if f.severity == "error"]
    return findings


def main(argv):
    if not argv or argv[0] in ("-h", "--help"):
        print(__doc__)
        return 2
    src = Path(argv[0])
    if not src.is_file():
        print(f"✗ 文件不存在: {src}")
        return 2
    level = "warning"
    json_out = None
    rest = argv[1:]
    if "--level" in rest:
        i = rest.index("--level")
        level = rest[i + 1]
        del rest[i:i + 2]
    if "--json" in rest:
        i = rest.index("--json")
        json_out = rest[i + 1]
        del rest[i:i + 2]
    try:
        prs = Presentation(str(src))
    except Exception as e:
        print(f"✗ 无法打开: {e}")
        return 2
    findings = audit(prs, level=level)
    if json_out:
        Path(json_out).write_text(
            json.dumps([asdict(f) for f in findings], ensure_ascii=False, indent=1),
            encoding="utf-8")
    n_err = sum(1 for f in findings if f.severity == "error")
    n_warn = len(findings) - n_err
    for f in findings:
        print(f"[{f.severity}] p{f.slide} {f.rule} shape={f.shape_id}: {f.message}")
    print(f"—— 共 {n_err} error / {n_warn} warning")
    return 1 if n_err else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
